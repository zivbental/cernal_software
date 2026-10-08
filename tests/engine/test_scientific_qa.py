"""Desired scientific behavior for the October 2026 audit regressions."""

import gc
import io
import weakref

import pytest
from openpyxl import Workbook

from engine.client import (
    LocalEngine,
    inspect_expression_input,
    normalize_trigger_sequence,
    validate_job_configuration,
)
from engine.domain import Constraints, DgeRow, DgeTable
from engine.errors import InputValidationError
from engine.gates.tools.folding import FoldEngine
from engine.inputs import parse_dge_table
from engine.pipeline import _build_constraints
from engine.scoring.normalize import failed_filter, normalize_value
from engine.scoring.profiles import DEFAULT_V1
from engine.stages.folding import FoldProfiler
from engine.stages.genes import GeneSelector
from engine.stages.motifs import MotifScreener


@pytest.mark.parametrize("suffix", ["csv", "tsv", "txt", "xlsx"])
@pytest.mark.parametrize("bad", ["inf", "-Infinity", "1e999", "not-a-number", "-0.1", "1.01"])
def test_invalid_probability_rejected_across_input_formats(suffix, bad):
    if suffix == "xlsx":
        book = Workbook()
        book.active.append(["gene_id", "log2fc", "pvalue"])
        book.active.append(["g1", 2.0, bad])
        stream = io.BytesIO()
        book.save(stream)
        raw = stream.getvalue()
    else:
        delimiter = "," if suffix == "csv" else "\t"
        raw = (
            delimiter.join(["gene_id", "log2fc", "pvalue"])
            + "\n"
            + delimiter.join(["g1", "2", bad])
            + "\n"
        ).encode()
    with pytest.raises(InputValidationError):
        parse_dge_table(raw, f"input.{suffix}")


@pytest.mark.parametrize("headers", ["gene_id,log2fc, log2 Fold Change", "gene_id,,log2fc"])
def test_normalized_header_collisions_rejected(headers):
    with pytest.raises(InputValidationError, match="headers"):
        parse_dge_table(f"{headers}\ng1,2,2\n".encode())


def test_preview_and_engine_use_first_xlsx_sheet(tmp_path):
    book = Workbook()
    book.active.append(["gene_id", "log2 Fold Change"])
    book.active.append(["first", 2])
    other = book.create_sheet("active-but-not-first")
    other.append(["gene_id", "log2fc"])
    other.append(["other", 9])
    book.active = 1
    path = tmp_path / "input.xlsx"
    book.save(path)
    report = inspect_expression_input(str(path))
    assert report["valid"]
    assert report["selected_sheet"] == 0
    assert report["preview"] == [{"gene_id": "first", "log2fc": 2.0}]
    assert parse_dge_table(path.read_bytes(), path.name).rows[0].gene_id == "first"


def test_raw_pvalues_need_complete_hypothesis_declaration():
    rows = tuple(DgeRow(f"g{i}", 2.0, p_value=p) for i, p in enumerate([0.01, 0.04, 0.9]))
    selector = GeneSelector(Constraints(direction_balance=False), MotifScreener())
    warnings = []
    subset = selector.select(
        DgeTable(rows, hypothesis_universe_complete=False), on_warning=warnings.append
    )
    assert len(subset) == 2
    assert all(g.p_adj is None for g in subset)
    assert any("not an FDR-controlled" in warning for warning in warnings)
    complete = selector.select(DgeTable(rows, hypothesis_universe_complete=True))
    assert len(complete) == 1
    assert complete[0].p_adj == pytest.approx(0.03)


def test_instance_fold_cache_is_bounded_and_collectible():
    folder = FoldEngine(cache_size=1)
    folder.mfe("GGGAAACCC")
    expected = folder.mfe("GGGAAACCC")
    folder.mfe("GGGAAAACCC")
    assert folder.mfe.cache_info().currsize == 1
    assert folder.mfe.cache_info().maxsize == 1
    assert folder.mfe("GGGAAACCC") == expected
    reference = weakref.ref(folder)
    del folder
    gc.collect()
    assert reference() is None


def test_instance_rnaplfold_cache_is_bounded_and_collectible():
    profiler = FoldProfiler()
    for sequence in ("GGGAAACCC", "GGGAAAACCC", "GGGAAAAACCC"):
        profiler.profile(sequence)
    assert profiler._matrix.cache_info().currsize == 2
    reference = weakref.ref(profiler)
    del profiler
    gc.collect()
    assert reference() is None


@pytest.mark.parametrize("bad", [float("nan"), float("inf"), -float("inf")])
def test_nonfinite_metrics_cannot_score_or_pass_filter(bad):
    with pytest.raises(ValueError, match="finite"):
        normalize_value(bad, DEFAULT_V1.spec("predicted_leakage"))
    assert failed_filter({"predicted_leakage": bad}, DEFAULT_V1) is not None


@pytest.mark.parametrize(
    "field,value",
    [
        ("max_genes", True),
        ("max_triggers", 0),
        ("max_p_adj", 2),
        ("min_separation", float("nan")),
        ("trigger_gc_range", [70, 30]),
        ("forbidden_motifs", [""]),
        ("forbidden_motifs", "ATG"),
        ("direction_balance", "false"),
    ],
)
def test_invalid_constraints_fail_before_computation(field, value):
    with pytest.raises(InputValidationError):
        _build_constraints({"constraints": {field: value}})


def test_forbidden_motifs_normalize_dna_case():
    assert _build_constraints({"constraints": {"forbidden_motifs": ["augc"]}}).forbidden_motifs == (
        "ATGC",
    )


@pytest.mark.parametrize("sequence", ["AUN", ">a\nAU\n>b\nGC", ">\nAUG", "AUG!"])
def test_trigger_normalization_never_discards_invalid_symbols(sequence):
    with pytest.raises(ValueError):
        normalize_trigger_sequence(sequence)


def test_single_fasta_record_normalizes_without_changing_coordinates():
    assert normalize_trigger_sequence(">one record\nat gc\nTT") == "AUGCUU"


def test_capabilities_advertise_only_runnable_host_family_pairs():
    capabilities = LocalEngine().capabilities()
    assert "antisense" not in capabilities.available_families
    assert all("and" not in name for name in capabilities.available_families)
    assert set(capabilities.family_hosts) == set(capabilities.available_families)
    for name, hosts in capabilities.family_hosts.items():
        for host in hosts:
            assert (
                validate_job_configuration({}, [name], "default", "direct", "AUGC", host)["host"]
                == host
            )
    with pytest.raises(ValueError, match="production supported"):
        validate_job_configuration(
            {}, ["prokaryotic_toehold"], "default", "direct", "AUGC", "human"
        )
